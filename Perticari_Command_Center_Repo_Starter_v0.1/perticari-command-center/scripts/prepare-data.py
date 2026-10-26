"""Create a private import package. Never writes to Mercado Libre or commits source data."""
import argparse, json, hashlib, re
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from collections import Counter
import openpyxl
import xlrd
p=argparse.ArgumentParser()
p.add_argument('--master',required=True);p.add_argument('--prices',required=True);p.add_argument('--technical',required=True);p.add_argument('--out',required=True)
a=p.parse_args()
def code(v): return str(v or '').strip()
def amount(v):
    if v is None or v=='': return None
    n=Decimal(str(v))*100
    return int(n.quantize(Decimal('1'),rounding=ROUND_HALF_UP)) if n>0 else None
book=xlrd.open_workbook(a.master);sheet=book.sheet_by_index(0);headers=sheet.row_values(0)
raw=[dict(zip(headers,sheet.row_values(i))) for i in range(1,sheet.nrows)]
counts=Counter(code(r['d6prod']) for r in raw)
products={}
for row in raw:
    sku=code(row['d6prod'])
    if not sku:continue
    name=str(row.get('d6desc') or '').strip()
    if not name:name='Descripción pendiente · '+sku
    products.setdefault(sku,{'sku':sku,'name':name,'base_cents':None,'currency':'ARS','price_confirmed':False,'source_at':None,'source':'archivo_MEMOSOFT','technical':{}})
    if counts[sku]>1:products[sku]['technical']['conflicto_origen']='Código repetido en el maestro; revisar equivalencia.'
    for field,label in [('d6des2','Descripción adicional'),('d6dete','Detalle de origen')]:
        if row.get(field):products[sku]['technical'][label]=str(row[field])
prices=openpyxl.load_workbook(a.prices,read_only=True,data_only=True).active
rows=prices.iter_rows(values_only=True);ph=next(rows)
for row in rows:
    record=dict(zip(ph,row));sku=code(record.get('Código'))
    if not sku:continue
    if sku not in products:products[sku]={'sku':sku,'name':str(record.get('Producto') or sku),'technical':{},'currency':'ARS','price_confirmed':False,'source_at':None,'source':'archivo_MEMOSOFT'}
    elif record.get('Producto'):products[sku]['name']=str(record['Producto']).strip()
    products[sku]['base_cents']=amount(record.get('Precio Original'))
listings={}
w=openpyxl.load_workbook(a.technical,read_only=True,data_only=True)
for ws in w:
    values=ws.iter_rows(values_only=True);first=next(values)
    if 'ID' not in first or 'SKU' not in first:continue
    columns=list(first);values=ws.iter_rows(min_row=5,values_only=True)
    for values_row in values:
        row=dict(zip(columns,values_row));item_id=code(row.get('ID'))
        if not re.fullmatch(r'MLA\d+',item_id):continue
        sku=code(row.get('SKU'));attributes={k:v for k,v in row.items() if k not in ['ID','SKU','FAMILY_ID','PRODUCT_NUMBER','VARIATION_ID','DOMAIN_ID','CATEGORY_ID','PARENT_CATEGORY_ID','TITLE'] and v not in [None,'']}
        if item_id not in listings:listings[item_id]={'id':item_id,'sku':sku,'title':str(row.get('TITLE') or ''),'status':'importado','price_cents':None,'currency':'ARS','category_id':code(row.get('CATEGORY_ID')),'source':'fichas_ML','source_at':1791849600000,'payload':{'product_number':code(row.get('PRODUCT_NUMBER')),'domain_id':code(row.get('DOMAIN_ID')),'group_id':code(row.get('FAMILY_ID')),'technical':attributes,'variations':[]}}
        if row.get('VARIATION_ID'):listings[item_id]['payload']['variations'].append({'id':code(row['VARIATION_ID']),'sku':sku,'attributes':attributes})
        if sku in products:products[sku]['technical'].setdefault('Ficha ML',attributes)
data={'schema':'perticari-private-snapshot-v1','source_files':{Path(path).name:hashlib.sha256(Path(path).read_bytes()).hexdigest() for path in [a.master,a.prices,a.technical]},'products':list(products.values()),'listings':list(listings.values())}
Path(a.out).parent.mkdir(parents=True,exist_ok=True);Path(a.out).write_text(json.dumps(data,ensure_ascii=False,default=str),encoding='utf-8')
print(json.dumps({'products':len(products),'listings':len(listings),'with_price':sum(v.get('base_cents') is not None for v in products.values()),'out':a.out},ensure_ascii=False))
